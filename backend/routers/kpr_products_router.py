"""Master Produk KPR per bank: tenor yang tersedia & suku bunga — sumber tunggal untuk form
pengajuan KPR (Customer & Kontrak) dan tahap SP3K (kontrak), menggantikan ketik bebas.
Termasuk impor massal dari Excel (template diunduh dari sini juga)."""
import io

from fastapi import APIRouter, Depends, File, HTTPException, Response, UploadFile
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from pydantic import BaseModel, Field

from db import db, ORG_ID
from core_utils import new_id, now_iso, serialize_doc
from rbac import require_permission, audit_log

router = APIRouter(prefix="/master/kpr-products", tags=["master-data"])
COLL = "kpr_products"

IMPORT_COLUMNS = [
    ("bank_name", "Bank (kode Kamus Data, mis. BTN/BNI/BRI/Mandiri/BCA) — wajib"),
    ("name", "Nama produk — wajib, unik per bank"),
    ("tenors", "Tenor bulan, pisahkan koma (mis. 120,180,240) — wajib"),
    ("interest_rate_pct", "Bunga %/tahun — wajib"),
    ("fixed_years", "Masa fixed (tahun) — opsional"),
    ("floating_rate_pct", "Bunga floating %/tahun — opsional"),
    ("min_dp_pct", "Min. DP % — opsional"),
    ("notes", "Catatan — opsional"),
    ("is_active", "Aktif? ya/tidak (kosong = ya)"),
]


class KprProductIn(BaseModel):
    bank_name: str = Field(min_length=1)
    name: str = Field(min_length=1)
    tenors: list[int] = Field(min_length=1)
    interest_rate_pct: float = Field(ge=0, le=100)
    fixed_years: int | None = Field(default=None, ge=0, le=40)
    floating_rate_pct: float | None = Field(default=None, ge=0, le=100)
    min_dp_pct: float | None = Field(default=None, ge=0, le=100)
    notes: str | None = None
    is_active: bool = True


def _clean(p: KprProductIn) -> dict:
    tenors = sorted({int(t) for t in p.tenors if 1 <= int(t) <= 480})
    if not tenors:
        raise HTTPException(status_code=400, detail="Isi minimal satu tenor (1–480 bulan).")
    d = p.model_dump()
    d.update({"bank_name": p.bank_name.strip(), "name": p.name.strip(), "tenors": tenors})
    return d


@router.get("")
async def list_products(bank_name: str = None, include_inactive: bool = False,
                        user: dict = Depends(require_permission("financing", "view"))):
    q = {"org_id": user.get("org_id", ORG_ID)}
    if bank_name:
        q["bank_name"] = bank_name
    if not include_inactive:
        q["is_active"] = True
    rows = await db[COLL].find(q, {"_id": 0}).sort([("bank_name", 1), ("name", 1)]).to_list(500)
    return {"data": serialize_doc(rows), "total": len(rows)}


@router.get("/import-template.xlsx")
async def import_template(user: dict = Depends(require_permission("financing", "view"))):
    wb = Workbook()
    ws = wb.active
    ws.title = "PRODUK_KPR"
    fill = PatternFill("solid", fgColor="1F3A5F")
    for c, (key, label) in enumerate(IMPORT_COLUMNS, start=1):
        ws.cell(row=1, column=c, value=key).font = Font(bold=True, color="FFFFFF")
        ws.cell(row=1, column=c).fill = fill
        ws.cell(row=2, column=c, value=label).font = Font(italic=True, color="666666")
        ws.column_dimensions[ws.cell(row=1, column=c).column_letter].width = 24 if key != "notes" else 40
    for r, ex in enumerate([
        ["BTN", "KPR Subsidi FLPP", "120,180,240", 5, 20, None, 1, "Rumah subsidi", "ya"],
        ["BNI", "BNI Griya Fixed 3", "60,120,180,240,300", 6.75, 3, 11.5, 10, "", "ya"],
    ], start=3):
        for c, v in enumerate(ex, start=1):
            ws.cell(row=r, column=c, value=v)
    ws.freeze_panes = "A3"
    info = wb.create_sheet("PETUNJUK")
    for i, t in enumerate([
        "Baris 1 = kunci kolom (jangan diubah), baris 2 = keterangan, baris 3 dst = data (contoh boleh dihapus/ditimpa).",
        "Produk dengan bank + nama yang SAMA dengan yang sudah ada akan DIPERBARUI, selain itu dibuat baru.",
        "Bank dianjurkan memakai kode Kamus Data (financing_bank) supaya cocok dengan pilihan bank di form pengajuan.",
        "Tenor dalam bulan (1–480), pisahkan dengan koma. Bunga & DP dalam persen tanpa tanda %.",
    ], start=1):
        info.cell(row=i, column=1, value=t)
    info.column_dimensions["A"].width = 120
    buf = io.BytesIO()
    wb.save(buf)
    return Response(buf.getvalue(),
                    media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": 'attachment; filename="SIPRO_Template_Produk_KPR.xlsx"'})


def _num(v, name, rn, errors, required=False):
    if v in (None, ""):
        if required:
            errors.append(f"Baris {rn}: {name} wajib diisi.")
        return None
    try:
        return float(str(v).replace("%", "").replace(",", ".").strip())
    except ValueError:
        errors.append(f"Baris {rn}: {name} bukan angka.")
        return None


def _parse_rows(content: bytes) -> tuple[list, list]:
    try:
        wb = load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    except Exception:  # noqa: BLE001
        raise HTTPException(status_code=400, detail="Berkas bukan Excel .xlsx yang sah.")
    ws = wb["PRODUK_KPR"] if "PRODUK_KPR" in wb.sheetnames else wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        raise HTTPException(status_code=400, detail="Sheet kosong.")
    header = [str(h or "").strip().lower() for h in rows[0]]
    if "bank_name" not in header or "name" not in header:
        raise HTTPException(status_code=400, detail="Baris 1 harus memuat kunci kolom template (bank_name, name, …).")
    idx = {k: header.index(k) for k, _ in IMPORT_COLUMNS if k in header}
    items, errors = [], []
    for rn, row in enumerate(rows[1:], start=2):
        get = lambda k: (row[idx[k]] if k in idx and idx[k] < len(row) else None)  # noqa: E731
        if all(v in (None, "") for v in row):
            continue
        if rn == 2 and str(get("bank_name") or "").lower().startswith("bank ("):
            continue
        bank, name = str(get("bank_name") or "").strip(), str(get("name") or "").strip()
        if not bank or not name:
            errors.append(f"Baris {rn}: bank dan nama produk wajib diisi.")
            continue
        tenors = sorted({int(float(t)) for t in str(get("tenors") or "").replace(";", ",").split(",")
                         if str(t).strip().replace(".", "").isdigit() and 1 <= int(float(t)) <= 480})
        if not tenors:
            errors.append(f"Baris {rn}: tenor kosong/tidak valid (bulan 1–480, pisahkan koma).")
            continue
        n_err = len(errors)
        rate = _num(get("interest_rate_pct"), "bunga", rn, errors, required=True)
        fixed = _num(get("fixed_years"), "masa fixed", rn, errors)
        floating = _num(get("floating_rate_pct"), "bunga floating", rn, errors)
        min_dp = _num(get("min_dp_pct"), "min. DP", rn, errors)
        if len(errors) > n_err:
            continue
        if not 0 <= rate <= 100:
            errors.append(f"Baris {rn}: bunga harus 0–100.")
            continue
        active = str(get("is_active") or "ya").strip().lower() not in ("tidak", "no", "false", "0", "n")
        items.append({"bank_name": bank, "name": name, "tenors": tenors, "interest_rate_pct": rate,
                      "fixed_years": int(fixed) if fixed is not None else None,
                      "floating_rate_pct": floating, "min_dp_pct": min_dp,
                      "notes": (str(get("notes")).strip() or None) if get("notes") not in (None, "") else None,
                      "is_active": active, "_row": rn})
    return items, errors


@router.post("/import")
async def import_products(file: UploadFile = File(...),
                          user: dict = Depends(require_permission("settings", "manage"))):
    """Impor massal: baris valid disimpan (bank+nama sama → diperbarui), baris salah dilaporkan."""
    if not (file.filename or "").lower().endswith(".xlsx"):
        raise HTTPException(status_code=400, detail="Berkas harus berekstensi .xlsx.")
    content = await file.read()
    if not content or len(content) > 5 * 1024 * 1024:
        raise HTTPException(status_code=400, detail="Berkas kosong atau melebihi 5 MB.")
    items, errors = _parse_rows(content)
    org = user.get("org_id", ORG_ID)
    created, updated = [], []
    seen = set()
    for it in items:
        rn = it.pop("_row")
        key = (it["bank_name"].lower(), it["name"].lower())
        if key in seen:
            errors.append(f"Baris {rn}: duplikat {it['bank_name']} / {it['name']} di berkas — dilewati.")
            continue
        seen.add(key)
        cur = await db[COLL].find_one({"org_id": org, "bank_name": it["bank_name"], "name": it["name"]}, {"_id": 0, "id": 1})
        if cur:
            await db[COLL].update_one({"id": cur["id"]}, {"$set": {**it, "updated_at": now_iso()}})
            updated.append(f"{it['bank_name']} / {it['name']}")
        else:
            doc = {"id": new_id(), "org_id": org, **it, "created_by": user.get("email"),
                   "created_at": now_iso(), "updated_at": now_iso()}
            await db[COLL].insert_one(doc)
            created.append(f"{it['bank_name']} / {it['name']}")
    await audit_log(user, "import", COLL, None, {"created": len(created), "updated": len(updated),
                                                  "errors": len(errors), "file": file.filename})
    return {"data": {"created": created, "updated": updated, "errors": errors,
                     "rows": len(created) + len(updated)}}


@router.post("")
async def create_product(payload: KprProductIn,
                         user: dict = Depends(require_permission("settings", "manage"))):
    org = user.get("org_id", ORG_ID)
    data = _clean(payload)
    if await db[COLL].find_one({"org_id": org, "bank_name": data["bank_name"], "name": data["name"]}):
        raise HTTPException(status_code=409, detail=f"Produk '{data['name']}' untuk {data['bank_name']} sudah ada.")
    doc = {"id": new_id(), "org_id": org, **data, "created_by": user.get("email"),
           "created_at": now_iso(), "updated_at": now_iso()}
    await db[COLL].insert_one(dict(doc))
    await audit_log(user, "create", COLL, doc["id"], {"bank": data["bank_name"], "name": data["name"]})
    doc.pop("_id", None)
    return {"data": serialize_doc(doc)}


@router.put("/{pid}")
async def update_product(pid: str, payload: KprProductIn,
                         user: dict = Depends(require_permission("settings", "manage"))):
    org = user.get("org_id", ORG_ID)
    data = _clean(payload)
    dup = await db[COLL].find_one({"org_id": org, "bank_name": data["bank_name"], "name": data["name"],
                                   "id": {"$ne": pid}})
    if dup:
        raise HTTPException(status_code=409, detail=f"Produk '{data['name']}' untuk {data['bank_name']} sudah ada.")
    res = await db[COLL].update_one({"id": pid, "org_id": org}, {"$set": {**data, "updated_at": now_iso()}})
    if not res.matched_count:
        raise HTTPException(status_code=404, detail="Produk KPR tidak ditemukan.")
    await audit_log(user, "update", COLL, pid, {"fields": sorted(data)})
    return {"data": serialize_doc(await db[COLL].find_one({"id": pid}, {"_id": 0}))}


@router.delete("/{pid}")
async def archive_product(pid: str, user: dict = Depends(require_permission("settings", "manage"))):
    """Arsip (nonaktif) — pengajuan yang sudah memakai produk ini tetap menyimpan angkanya."""
    org = user.get("org_id", ORG_ID)
    res = await db[COLL].update_one({"id": pid, "org_id": org},
                                    {"$set": {"is_active": False, "updated_at": now_iso()}})
    if not res.matched_count:
        raise HTTPException(status_code=404, detail="Produk KPR tidak ditemukan.")
    await audit_log(user, "archive", COLL, pid)
    return {"data": {"id": pid, "is_active": False}}
