"""Iteration 26 backend tests:
- Excel import template download (GET /api/master/kpr-products/import-template.xlsx)
- Excel import upload (POST /api/master/kpr-products/import) — valid, template-roundtrip,
  errors (missing name, bad tenors, bad rate, duplicate in file, inactive), RBAC (sales 403).
- Docgen tokens: kpr_terms / kpr_terms_line / dp_line contain 'Fasilitas KPR :' for KPR
  scheme. Requires the seeded demo contract to be advanced to SP3K first.
- Regression: list still works, POST duplicate 409.

Uses env REACT_APP_BACKEND_URL for HTTP; also loads /app/backend/.env for in-process import.
"""
import io
import os
import sys
import time
import uuid

import pytest
import requests
from openpyxl import Workbook, load_workbook

BASE = os.environ["REACT_APP_BACKEND_URL"].rstrip("/")
API = f"{BASE}/api"
SUPER = ("superadmin@sipro.co.id", "Sipro#2026")
SALES = ("sales@sipro.co.id", "Sipro#2026")


def _login(email, pw):
    r = requests.post(f"{API}/auth/login", json={"email": email, "password": pw}, timeout=30)
    assert r.status_code == 200, r.text
    return r.json()["access_token"]


@pytest.fixture(scope="module")
def super_h():
    return {"Authorization": f"Bearer {_login(*SUPER)}"}


@pytest.fixture(scope="module")
def sales_h():
    return {"Authorization": f"Bearer {_login(*SALES)}"}


# ------------------------------------------------------------------ Template
class TestTemplate:
    def test_download_template_super(self, super_h):
        r = requests.get(f"{API}/master/kpr-products/import-template.xlsx", headers=super_h, timeout=30)
        assert r.status_code == 200, r.text
        ct = r.headers.get("content-type", "")
        assert "spreadsheetml" in ct, ct
        cd = r.headers.get("content-disposition", "")
        assert "SIPRO_Template_Produk_KPR.xlsx" in cd, cd
        wb = load_workbook(io.BytesIO(r.content), read_only=True)
        assert "PRODUK_KPR" in wb.sheetnames
        ws = wb["PRODUK_KPR"]
        header = [str(c.value or "").strip() for c in ws[1]]
        for k in ("bank_name", "name", "tenors", "interest_rate_pct", "fixed_years",
                  "floating_rate_pct", "min_dp_pct", "notes", "is_active"):
            assert k in header, f"missing {k}"

    def test_download_template_sales_ok(self, sales_h):
        r = requests.get(f"{API}/master/kpr-products/import-template.xlsx", headers=sales_h, timeout=30)
        assert r.status_code == 200  # financing:view permits it


# ------------------------------------------------------------------ Import
def _make_xlsx(rows: list[list]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "PRODUK_KPR"
    header = ["bank_name", "name", "tenors", "interest_rate_pct", "fixed_years",
              "floating_rate_pct", "min_dp_pct", "notes", "is_active"]
    ws.append(header)
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


class TestImport:
    def test_import_template_roundtrip(self, super_h):
        r = requests.get(f"{API}/master/kpr-products/import-template.xlsx", headers=super_h, timeout=30)
        files = {"file": ("template.xlsx", r.content,
                          "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
        r = requests.post(f"{API}/master/kpr-products/import", headers=super_h, files=files, timeout=30)
        assert r.status_code == 200, r.text
        data = r.json()["data"]
        # Template has 2 example rows; both should be created or updated
        assert data["rows"] == 2, data
        # BTN/KPR Subsidi FLPP is pre-seeded — should be updated
        names = [x for x in data["created"] + data["updated"]]
        assert any("BTN" in n and "KPR Subsidi FLPP" in n for n in names), names

    def test_import_custom_mix(self, super_h):
        ts = int(time.time())
        good_name = f"Mandiri KPR Test {ts}"
        content = _make_xlsx([
            ["Mandiri", good_name, "120,240", 7.5, None, None, None, "ok", "ya"],
            ["Mandiri", "", "120", 7.5, None, None, None, "missing name", "ya"],  # error missing name
            ["Mandiri", f"BadTenors {ts}", "abc", 7.5, None, None, None, "", "ya"],  # error tenors
            ["Mandiri", f"BadRate {ts}", "120", "xx", None, None, None, "", "ya"],  # error rate
            ["Mandiri", good_name, "120,240", 7.5, None, None, None, "dup", "ya"],  # duplicate in file
            ["Mandiri", f"Inactive {ts}", "120", 7.5, None, None, None, "off", "tidak"],  # created inactive
        ])
        files = {"file": ("mix.xlsx", content,
                          "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
        r = requests.post(f"{API}/master/kpr-products/import", headers=super_h, files=files, timeout=30)
        assert r.status_code == 200, r.text
        d = r.json()["data"]
        # Expect 2 successful rows: the good one + inactive one
        assert d["rows"] == 2, d
        errs = " | ".join(d["errors"])
        assert "Baris 3" in errs and "wajib" in errs.lower(), errs   # missing name
        assert "Baris 4" in errs, errs                                # bad tenors
        assert "Baris 5" in errs, errs                                # bad rate
        assert "Baris 6" in errs and "duplikat" in errs.lower(), errs # duplicate in file
        # inactive product must exist and only visible with include_inactive
        r2 = requests.get(f"{API}/master/kpr-products",
                          params={"bank_name": "Mandiri"}, headers=super_h, timeout=30)
        names_active = [p["name"] for p in r2.json()["data"]]
        assert good_name in names_active
        assert f"Inactive {ts}" not in names_active
        r3 = requests.get(f"{API}/master/kpr-products",
                          params={"bank_name": "Mandiri", "include_inactive": "true"},
                          headers=super_h, timeout=30)
        inactive_row = [p for p in r3.json()["data"] if p["name"] == f"Inactive {ts}"]
        assert inactive_row and inactive_row[0]["is_active"] is False

        # Cleanup: archive created Mandiri test products
        for p in r3.json()["data"]:
            if p["name"].startswith(("Mandiri KPR Test ", "Inactive ")) and str(ts) in p["name"]:
                requests.delete(f"{API}/master/kpr-products/{p['id']}",
                                headers=super_h, timeout=30)

    def test_import_rejects_non_xlsx(self, super_h):
        files = {"file": ("bad.csv", b"bank_name,name\nBTN,x\n", "text/csv")}
        r = requests.post(f"{API}/master/kpr-products/import", headers=super_h, files=files, timeout=30)
        assert r.status_code == 400, r.text

    def test_sales_cannot_import(self, sales_h):
        content = _make_xlsx([["BTN", f"NoPerm {uuid.uuid4().hex[:6]}", "120", 5, None, None, None, "", "ya"]])
        files = {"file": ("x.xlsx", content,
                          "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
        r = requests.post(f"{API}/master/kpr-products/import", headers=sales_h, files=files, timeout=30)
        assert r.status_code == 403, r.text


# ------------------------------------------------------------------ Regression
class TestRegression:
    def test_list_still_works(self, super_h):
        r = requests.get(f"{API}/master/kpr-products", headers=super_h, timeout=30)
        assert r.status_code == 200
        assert "data" in r.json()

    def test_duplicate_post_409(self, super_h):
        # BTN/KPR Subsidi FLPP always exists after template roundtrip
        r = requests.post(f"{API}/master/kpr-products", headers=super_h, timeout=30, json={
            "bank_name": "BTN", "name": "KPR Subsidi FLPP", "tenors": [120], "interest_rate_pct": 5,
        })
        assert r.status_code == 409, r.text


# ------------------------------------------------------------------ SP3K + docgen tokens
DEMO_CONTRACT = "c3e19de7-1358-4049-a74b-d062cdedfb01"
DEMO_CUSTOMER = "6e481f66-f594-4088-8f50-38d82812fd7e"


def _advance_to_sp3k_if_needed(super_h) -> str:
    """Return contract_id currently sitting at SP3K gate. Reuse demo, else reseed."""
    r = requests.get(f"{API}/contracts/{DEMO_CONTRACT}/kpr", headers=super_h, timeout=30)
    if r.status_code == 200 and r.json().get("data", {}).get("next_stage") == "sp3k":
        return DEMO_CONTRACT
    # try re-seeding via script (imports over HTTP, so subprocess is fine)
    import subprocess
    p = subprocess.run(["python3", "/app/scripts/seed_kpr_sp3k_demo.py"],
                       capture_output=True, text=True, timeout=90)
    out = (p.stdout or "") + (p.stderr or "")
    for line in out.splitlines():
        if line.startswith("contract_id"):
            return line.split(":", 1)[1].strip()
    pytest.skip(f"cannot ensure sp3k demo: {out[:400]}")


class TestSp3kAndDocgen:
    def test_sp3k_record_updates_application(self, super_h):
        cid = _advance_to_sp3k_if_needed(super_h)
        # upload an evidence PDF
        up = requests.post(f"{API}/files/upload", headers=super_h, timeout=60,
                           files={"file": ("sp3k.pdf", b"%PDF-1.4 SP3K demo\n", "application/pdf")},
                           data={"owner_type": "contract", "owner_id": cid, "optimize": "false"})
        assert up.status_code == 200, up.text
        fid = up.json()["data"]["id"]

        # find BTN product id
        r = requests.get(f"{API}/master/kpr-products", params={"bank_name": "BTN"},
                         headers=super_h, timeout=30)
        prods = [p for p in r.json()["data"] if p["name"] == "KPR Subsidi FLPP"]
        assert prods, "seed produk KPR BTN missing"
        prod = prods[0]

        payload = {
            "number": f"SP3K-TEST-{int(time.time())}",
            "plafon": 250_000_000,
            "tenor_months": 240,
            "rate": 5,
            "kpr_product_id": prod["id"],
            "kpr_product_name": prod["name"],
            "valid_until": "2026-12-31",
            "file_id": fid,
            "note": "iter26 pytest",
        }
        r = requests.post(f"{API}/contracts/{cid}/kpr/stage/sp3k", headers=super_h,
                          json=payload, timeout=30)
        assert r.status_code == 200, r.text
        # verify application persisted the product fields
        r = requests.get(f"{API}/contracts/{cid}/kpr", headers=super_h, timeout=30)
        app = r.json()["data"]["application"]
        assert app.get("kpr_product_id") == prod["id"]
        assert app.get("tenor_months") == 240
        assert float(app.get("interest_rate_pct") or 0) == 5
        # sp3k sub-object should carry product name
        sp3k = app.get("sp3k") or {}
        assert sp3k.get("product_name") == "KPR Subsidi FLPP" or app.get("kpr_product_name") == "KPR Subsidi FLPP"

    def test_docgen_tokens_kpr_terms(self, super_h):
        """Import docgen.build_context directly (backend .env preloaded)."""
        from dotenv import load_dotenv
        load_dotenv("/app/backend/.env")
        sys.path.insert(0, "/app/backend")
        import asyncio
        import docgen
        from db import db, ORG_ID

        async def _run():
            contract = await db.contracts.find_one({"id": DEMO_CONTRACT}, {"_id": 0})
            assert contract, "demo contract not found"
            ctx = await docgen.build_context(ORG_ID, contract, "SPR_KPR",
                                             actor_name="pytest", doc_number="TEST/SPR-KPR/DEMO")
            assert "BTN" in ctx["kpr_terms"], ctx["kpr_terms"]
            assert "KPR Subsidi FLPP" in ctx["kpr_terms"], ctx["kpr_terms"]
            assert "tenor 240 bulan" in ctx["kpr_terms"], ctx["kpr_terms"]
            assert "bunga 5%/th" in ctx["kpr_terms"], ctx["kpr_terms"]
            assert ctx["kpr_terms_line"].startswith("Fasilitas KPR :"), ctx["kpr_terms_line"]
            assert "Fasilitas KPR :" in ctx["dp_line"], ctx["dp_line"]

        asyncio.get_event_loop().run_until_complete(_run()) if not asyncio.get_event_loop().is_running() else asyncio.run(_run())
